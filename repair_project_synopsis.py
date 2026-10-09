from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
SOURCE = ROOT / 'PROJECT_SYNOPSIS.md'
TARGET = ROOT / 'PROJECT_SYNOPSIS.docx'
ORIGINAL = ROOT / 'PROJECT_SYNOPSIS.docx'

BLUE = RGBColor(30, 91, 168)
TEAL = RGBColor(0, 128, 128)
DARK = RGBColor(31, 41, 55)
MUTED = RGBColor(75, 85, 99)
LIGHT_BLUE = 'EAF3FF'
LIGHT_TEAL = 'E7F7F5'
BORDER = 'CBD5E1'


def cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        properties.append(shading)
    shading.set(qn('w:fill'), fill)


def set_cell_border(cell, color=BORDER, size='6'):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        properties.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Page ')
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(field)


def style_document(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [('Heading 1', 17, BLUE), ('Heading 2', 13, TEAL), ('Heading 3', 11.5, BLUE)]:
        style = doc.styles[name]
        style.font.name = 'Aptos Display'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    if 'Code Block' not in [s.name for s in doc.styles]:
        code = doc.styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = 'Consolas'
        code.font.size = Pt(8.5)
        code.font.color.rgb = DARK
        code.paragraph_format.left_indent = Inches(0.2)
        code.paragraph_format.space_after = Pt(0)
    if 'Caption Text' not in [s.name for s in doc.styles]:
        caption = doc.styles.add_style('Caption Text', WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = 'Aptos'
        caption.font.size = Pt(9)
        caption.font.italic = True
        caption.font.color.rgb = MUTED


def add_inline(paragraph, text, size=None):
    pattern = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')
    for part in pattern.split(text):
        if not part:
            continue
        run = paragraph.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
        run.font.name = 'Consolas' if part.startswith('`') else 'Aptos'
        run.bold = part.startswith('**')
        if size:
            run.font.size = Pt(size)


def add_text(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    add_inline(p, text)
    return p


def add_table(doc, rows):
    data = []
    for line in rows:
        values = [v.strip() for v in line.strip().strip('|').split('|')]
        if values and not all(set(v) <= set('-: ') for v in values):
            data.append(values)
    if not data:
        return
    table = doc.add_table(rows=1, cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, value in enumerate(data[0]):
        cell = table.rows[0].cells[i]
        cell.text = value.replace('**', '')
        cell_shading(cell, '1E5BA8')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.name = 'Aptos'
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
    for row_index, values in enumerate(data[1:]):
        cells = table.add_row().cells
        for i in range(len(data[0])):
            value = values[i] if i < len(values) else ''
            cells[i].text = value.replace('**', '')
            cell_shading(cells[i], LIGHT_BLUE if row_index % 2 == 0 else 'FFFFFF')
            set_cell_border(cells[i])
            for run in cells[i].paragraphs[0].runs:
                run.font.name = 'Aptos'
                run.font.size = Pt(9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def existing_edits():
    old = Document(ORIGINAL)
    texts = [p.text.strip() for p in old.paragraphs if p.text.strip()]
    objective = next((t.split(':', 1)[1].strip() for t in texts if t.startswith('Objective :')), None)
    overview = next((t for t in texts if t.startswith('IAMShield AI is an autonomous policy synthesis engine')), None)
    features = [t for t in texts if ':-' in t and t.endswith('|')]
    return objective, overview, features


def build():
    objective, edited_overview, edited_features = existing_edits()
    source_lines = SOURCE.read_text(encoding='utf-8').splitlines()
    doc = Document()
    style_document(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    header = section.header.paragraphs[0]
    header.text = 'IAMShield AI  |  Project Synopsis'
    header.runs[0].font.name = 'Aptos'
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = MUTED
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('IAMShield AI')
    run.font.name = 'Aptos Display'
    run.font.size = Pt(29)
    run.font.bold = True
    run.font.color.rgb = BLUE
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('Autonomous Least-Privilege IAM Policy Synthesizer')
    run.font.size = Pt(14)
    run.font.color.rgb = TEAL
    doc.add_paragraph()

    in_code = False
    code_lines = []
    list_items = []
    table_lines = []
    current_section = ''
    skipped_proposed = False

    def flush_lists():
        nonlocal list_items
        if list_items:
            for item in list_items:
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.left_indent = Inches(0.22)
                add_inline(p, item)
            list_items = []

    def flush_table():
        nonlocal table_lines
        if table_lines:
            add_table(doc, table_lines)
            table_lines = []

    def flush_code():
        nonlocal code_lines
        if code_lines:
            block = doc.add_table(rows=1, cols=1)
            cell = block.cell(0, 0)
            cell_shading(cell, 'F1F5F9')
            set_cell_border(cell, 'D5DEE8', '4')
            cell.text = ''
            for index, code in enumerate(code_lines):
                p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
                p.style = 'Code Block'
                p.add_run(code)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            code_lines = []

    for raw in source_lines:
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith('```'):
            flush_lists(); flush_table()
            if in_code:
                flush_code()
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if stripped.startswith('|'):
            flush_lists()
            table_lines.append(stripped)
            continue
        flush_table()
        if not stripped or stripped == '---':
            flush_lists()
            continue
        if stripped.startswith('# '):
            continue
        if stripped.startswith('## '):
            flush_lists()
            heading = stripped[3:]
            current_section = heading
            doc.add_heading(heading, level=1)
            if heading == 'Executive Summary':
                add_text(doc, 'Project Title: IAMShield AI: Autonomous Least-Privilege IAM Policy Synthesizer')
                add_text(doc, f'Objective: {objective or "To develop an intelligent system that autonomously analyzes access patterns and synthesizes least-privilege IAM policies."}')
                add_text(doc, 'Credits: 3-Credit Mini Project')
                add_text(doc, 'Status: Research & Design Phase')
                add_text(doc, 'Date: September 2, 2026')
            if heading == '2. Proposed Solution':
                skipped_proposed = False
            continue
        if stripped.startswith('### '):
            flush_lists()
            if current_section == '2. Proposed Solution' and stripped[4:] == 'System Overview':
                add_text(doc, edited_overview or 'IAMShield AI is an autonomous policy synthesis engine for cloud environments.')
                add_text(doc, 'It continuously analyzes runtime application behavior to synthesize, test, and enforce precise least-privilege Identity and Access Management policies with controlled human oversight.')
                add_text(doc, 'Key Features', style='Heading 3')
                for feature in edited_features:
                    add_text(doc, feature.rstrip('|').strip())
                skipped_proposed = True
                continue
            if current_section == '2. Proposed Solution' and skipped_proposed:
                continue
            doc.add_heading(stripped[4:], level=2)
            continue
        if stripped.startswith('#### '):
            flush_lists()
            doc.add_heading(stripped[5:], level=3)
            continue
        if current_section == '2. Proposed Solution' and skipped_proposed:
            continue
        if stripped.startswith('- ') or stripped.startswith('* ') or re.match(r'^\d+\. ', stripped):
            list_items.append(re.sub(r'^[-*]\s+', '', stripped))
            continue
        flush_lists()
        if stripped.startswith('**') and stripped.endswith('**'):
            add_text(doc, stripped.strip('*'), style='Heading 3')
        elif stripped.startswith('> '):
            p = add_text(doc, stripped[2:])
            p.paragraph_format.left_indent = Inches(0.25)
        else:
            add_text(doc, stripped)
    flush_lists(); flush_table(); flush_code()
    doc.save(TARGET)
    print(f'Repaired DOCX: {TARGET}')


if __name__ == '__main__':
    build()
